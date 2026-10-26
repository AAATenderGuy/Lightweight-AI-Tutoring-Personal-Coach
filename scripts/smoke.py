"""开发用冒烟测试：依次跑通资料、对话、课堂、测验、进度。

用法（服务已启动）：
    .venv/Scripts/python scripts/smoke.py http://127.0.0.1:8000
"""
from __future__ import annotations

import io
import json
import sys
import time
from pathlib import Path

import httpx

BASE = sys.argv[1].rstrip("/") if len(sys.argv) > 1 else "http://127.0.0.1:8000"
ROOT = Path(__file__).resolve().parent.parent


def show(title: str) -> None:
    print(f"\n=== {title} ===")


def read_sse(client: httpx.Client, path: str, payload: dict) -> dict:
    """收集 SSE 事件，返回 {"types": [...], "lesson": ..., "text": ...}。"""
    result: dict = {
        "types": [],
        "text": "",
        "lesson": None,
        "sources": [],
        "agents": [],
        "agent_starts": [],
    }
    with client.stream("POST", f"{BASE}{path}", json=payload, timeout=120) as response:
        response.raise_for_status()
        for line in response.iter_lines():
            if not line.startswith("data:"):
                continue
            event = json.loads(line[5:].strip())
            result["types"].append(event.get("type"))
            if event.get("type") == "delta":
                result["text"] += event.get("text", "")
            elif event.get("type") == "lesson":
                result["lesson"] = event.get("lesson")
            elif event.get("type") == "meta":
                result["sources"] = event.get("sources", [])
                result["agents"] = event.get("agents", [])
            elif event.get("type") == "agent_start":
                result["agent_starts"].append(event.get("agent"))
            elif event.get("type") == "error":
                raise RuntimeError(event.get("message"))
    return result


def main() -> None:
    with httpx.Client(timeout=60) as client:
        show("health")
        health = client.get(f"{BASE}/api/health").json()
        print(json.dumps(health, ensure_ascii=False))
        assert health["ok"] and health["mode"] in ("demo", "live")

        show("materials")
        sample = ROOT / "samples" / "光合作用.md"
        with sample.open("rb") as handle:
            uploaded = client.post(
                f"{BASE}/api/materials",
                files={"file": (sample.name, handle, "text/markdown")},
            )
        print(uploaded.status_code, uploaded.json().get("title"), uploaded.json().get("chunks"))
        uploaded.raise_for_status()
        material_id = uploaded.json()["id"]
        listing = client.get(f"{BASE}/api/materials").json()
        assert any(item["id"] == material_id for item in listing["items"])

        show("chat (socratic)")
        started = time.time()
        chat = read_sse(
            client,
            "/api/chat/stream",
            {
                "message": "光合作用的暗反应为什么在夜里会停止？",
                "mode": "socratic",
                "history": [],
                "material_ids": [],
                "use_materials": True,
            },
        )
        print(f"{len(chat['text'])} chars in {time.time() - started:.1f}s, sources={len(chat['sources'])}")
        print(chat["text"][:120].replace("\n", " "))
        assert chat["text"]

        show("classroom generate")
        started = time.time()
        classroom = read_sse(
            client,
            "/api/classroom/generate",
            {
                "topic": "光合作用",
                "level": "beginner",
                "duration": 20,
                "material_ids": [material_id],
                "scene_count": 6,
            },
        )
        lesson = classroom["lesson"]
        assert lesson and lesson["scenes"], "lesson is empty"
        print(f"{lesson['title']} · {len(lesson['scenes'])} scenes in {time.time() - started:.1f}s")
        for scene in lesson["scenes"]:
            print(" -", scene["type"], scene["title"])

        show("classroom discussion")
        first_content = next((scene for scene in lesson["scenes"] if scene["type"] == "content"), lesson["scenes"][0])
        talk = read_sse(
            client,
            "/api/classroom/discuss",
            {
                "lesson_id": lesson["id"],
                "scene_id": first_content["id"],
                "question": "这一页的核心概念是什么？",
                "history": [],
            },
        )
        print(talk["text"][:120])
        follow_up = read_sse(
            client,
            "/api/classroom/discuss",
            {
                "lesson_id": lesson["id"],
                "scene_id": first_content["id"],
                "question": "能不能再举一个更具体的例子？",
                "history": [
                    {"role": "user", "content": "这一页的核心概念是什么？"},
                    {"role": "assistant", "content": talk["text"]},
                ],
            },
        )
        print("follow-up:", follow_up["text"][:80])
        assert follow_up["text"], "teacher follow-up is empty"

        show("group interaction")
        group_agents = client.get(f"{BASE}/api/classroom/group/agents").json()["agents"]
        print("agents:", [agent["name"] for agent in group_agents])
        assert len(group_agents) >= 3
        auto_group = read_sse(
            client,
            "/api/classroom/group-discuss",
            {
                "lesson_id": lesson["id"],
                "scene_id": first_content["id"],
                "message": "",
                "mode": "auto",
                "history": [],
            },
        )
        print("auto speakers:", [agent["name"] for agent in auto_group["agent_starts"]])
        print(auto_group["text"][:120])
        assert len(auto_group["agent_starts"]) >= 2, "auto group did not gather classmates"
        assert auto_group["text"]

        group_chat = read_sse(
            client,
            "/api/classroom/group-discuss",
            {
                "lesson_id": lesson["id"],
                "scene_id": first_content["id"],
                "message": "我觉得光合作用只在白天发生，这个说法对吗？",
                "mode": "chat",
                "history": [
                    {"role": "assistant", "agent": "小问", "content": "我们先确认一下光合作用的条件。"},
                ],
            },
        )
        print("group reply:", group_chat["text"][:100])
        assert group_chat["text"]

        show("quiz")
        generated = client.post(
            f"{BASE}/api/quiz/generate",
            json={"topic": "光合作用", "material_ids": [material_id], "count": 3, "difficulty": "medium"},
        )
        print(generated.status_code)
        generated.raise_for_status()
        quiz = generated.json()["quiz"]
        print(quiz["title"], "questions:", len(quiz["questions"]))
        answers = []
        for question in quiz["questions"]:
            if question["type"] == "single":
                answers.append({"question_id": question["id"], "answer": "0"})
            else:
                answers.append({"question_id": question["id"], "answer": "光反应把光能转化成 ATP 和 NADPH，暗反应利用它们固定二氧化碳。"})
        graded = client.post(f"{BASE}/api/quiz/{quiz['id']}/submit", json={"answers": answers}, timeout=120)
        graded.raise_for_status()
        result = graded.json()
        print(f"score={result['score']} correct={result['correct_count']}/{result['total']}")
        for item in result["results"]:
            assert "correct_answer" in item and "explanation" in item
            print(" -", item["score"], item["feedback"][:50])

        show("quiz similar")
        focus = [item["knowledge_point"] for item in result["results"] if not item["correct"] and item.get("knowledge_point")]
        similar = client.post(
            f"{BASE}/api/quiz/{quiz['id']}/similar",
            json={"count": 3, "focus_points": focus},
            timeout=120,
        )
        similar.raise_for_status()
        practice = similar.json()["quiz"]
        assert practice["practice"] is True
        assert practice["questions"]
        for question in practice["questions"]:
            assert question.get("explanation"), "practice question has no explanation"
            if question["type"] == "single":
                assert "answer_index" in question
        print(practice["title"], "questions:", len(practice["questions"]), "practice:", practice["practice"])
        practice_answers = []
        for question in practice["questions"]:
            if question["type"] == "single":
                practice_answers.append({"question_id": question["id"], "answer": str(question["answer_index"])})
            else:
                practice_answers.append({"question_id": question["id"], "answer": "用自己的话说明并配一个例子。"})
        practiced = client.post(
            f"{BASE}/api/quiz/{practice['id']}/submit",
            json={"answers": practice_answers},
            timeout=120,
        )
        practiced.raise_for_status()
        print("similar score:", practiced.json()["score"])

        show("pptx export")
        exported = client.get(f"{BASE}/api/classroom/lessons/{lesson['id']}/export/pptx", timeout=60)
        exported.raise_for_status()
        assert exported.content[:2] == b"PK", "export is not a zip/pptx"
        assert "presentationml" in exported.headers.get("content-type", "")
        from pptx import Presentation

        deck = Presentation(io.BytesIO(exported.content))
        print(f"{len(exported.content) / 1024:.1f} KB, {len(deck.slides)} slides")
        assert len(deck.slides) >= len(lesson["scenes"])
        notes = deck.slides[1].notes_slide.notes_text_frame.text
        assert notes.strip(), "speaker notes are empty"
        print("speaker notes chars:", len(notes))

        show("progress")
        progress = client.get(f"{BASE}/api/progress").json()
        print(json.dumps(progress["stats"], ensure_ascii=False))
        for topic in progress["topics"]:
            print(f" - {topic['topic']}: {topic['mastery']} ({topic['label']}, {topic['attempts']} attempts)")

        show("lessons + settings")
        lessons = client.get(f"{BASE}/api/classroom/lessons").json()["items"]
        print(f"{len(lessons)} lessons saved")
        print(json.dumps(client.get(f"{BASE}/api/settings").json(), ensure_ascii=False))

    print("\nSMOKE OK")


if __name__ == "__main__":
    main()
