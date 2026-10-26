"""把课堂 JSON 导出为教师版 PPTX（含讲解词备注与答案附录）。"""
from __future__ import annotations

import io
import re

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

INK = RGBColor(0x17, 0x17, 0x1A)
INK2 = RGBColor(0x55, 0x55, 0x4F)
INK3 = RGBColor(0x86, 0x86, 0x7E)
BORDER = RGBColor(0xE5, 0xE5, 0xDF)
ACCENT = RGBColor(0x1F, 0x6B, 0x4C)
ACCENT_TEXT = RGBColor(0x17, 0x4D, 0x37)
ACCENT_SOFT = RGBColor(0xEA, 0xF2, 0xED)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Microsoft YaHei"
SLIDE_W = 13.333
SLIDE_H = 7.5

TYPE_LABELS = {
    "cover": "开场",
    "content": "讲解",
    "discussion": "讨论",
    "quiz": "检测",
    "summary": "回顾",
}


def _set_font(run, size: float, bold: bool = False, color: RGBColor = INK, font: str = FONT) -> None:
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font
    rpr = run._r.get_or_add_rPr()
    east_asian = rpr.find(qn("a:ea"))
    if east_asian is None:
        east_asian = etree.SubElement(rpr, qn("a:ea"))
    east_asian.set("typeface", font)


def _textbox(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    *,
    size: float = 18,
    bold: bool = False,
    color: RGBColor = INK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing: float = 1.25,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    _set_font(run, size, bold, color)


def _paragraphs(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    lines: list[str],
    *,
    size: float = 18,
    color: RGBColor = INK,
    bold: bool = False,
    bullet: bool = False,
    space_after: float = 9,
    line_spacing: float = 1.35,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, text in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.line_spacing = line_spacing
        paragraph.space_after = Pt(space_after)
        run = paragraph.add_run()
        run.text = f"•  {text}" if bullet else text
        _set_font(run, size, bold, color)


def _accent_bar(slide) -> None:
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W), Inches(0.1))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT
    bar.line.fill.background()


def _divider(slide, top: float, left: float = 0.9, width: float = 11.53) -> None:
    line = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(0.012)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = BORDER
    line.line.fill.background()


def _footer(slide, title: str, page: int) -> None:
    _divider(slide, 6.95)
    label = title if len(title) <= 46 else title[:46] + "…"
    _textbox(slide, 0.9, 7.03, 10.4, 0.3, f"Edu · {label}", size=9, color=INK3)
    _textbox(
        slide, 11.4, 7.03, 1.03, 0.3, str(page), size=9, color=INK3, align=PP_ALIGN.RIGHT
    )


def _header(slide, kicker: str, title: str, page: int, footer_title: str) -> None:
    _accent_bar(slide)
    _textbox(slide, 0.9, 0.48, 11.5, 0.3, kicker, size=10, color=INK3)
    _textbox(slide, 0.9, 0.88, 11.5, 1.0, title, size=28, bold=True, color=INK)
    _divider(slide, 2.02)
    _footer(slide, footer_title, page)


def _prompt_box(slide, text: str, top: float = 5.45) -> None:
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.9), Inches(top), Inches(11.53), Inches(0.9)
    )
    box.adjustments[0] = 0.08
    box.fill.solid()
    box.fill.fore_color.rgb = ACCENT_SOFT
    box.line.fill.background()
    box.shadow.inherit = False
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.22)
    frame.margin_right = Inches(0.22)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.04)
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    paragraph.line_spacing = 1.3
    run = paragraph.add_run()
    run.text = text
    _set_font(run, 14, False, ACCENT_TEXT)


def _note(slide, text: str) -> None:
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _bullets_size(bullets: list[str]) -> float:
    total = sum(len(item) for item in bullets)
    if total > 220:
        return 14
    if total > 150:
        return 16
    return 18


def _clean_bullets(value, limit: int = 5) -> list[str]:
    result = []
    for item in value or []:
        text = re.sub(r"\s+", " ", str(item)).strip()
        if text:
            result.append(text[:110])
        if len(result) >= limit:
            break
    return result


def _question_lines(question: dict, index: int) -> list[str]:
    lines = [f"{index}. {str(question.get('question') or '').strip()[:150]}"]
    options = [str(option).strip() for option in (question.get("options") or []) if str(option).strip()]
    for option_index, option in enumerate(options[:6]):
        lines.append(f"{'ABCDEF'[option_index]}. {option[:80]}")
    return lines


def _answer_lines(question: dict, index: int) -> list[str]:
    lines = [f"{index}. {str(question.get('question') or '').strip()[:150]}"]
    if question.get("type") == "single" and question.get("options"):
        try:
            correct = int(question.get("answer_index", 0))
        except (TypeError, ValueError):
            correct = 0
        correct = max(0, min(len(question["options"]) - 1, correct))
        lines.append(f"正确答案：{'ABCDEF'[correct]}. {question['options'][correct]}")
    else:
        lines.append(f"参考答案：{str(question.get('answer') or '（暂无）')[:220]}")
    explanation = str(question.get("explanation") or "").strip()
    if explanation:
        lines.append(f"解析：{explanation[:260]}")
    return lines


def _add_cover(prs, lesson: dict, page: int) -> int:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _accent_bar(slide)
    _textbox(slide, 0.9, 0.75, 11.5, 0.35, "EDU · 迷你课堂", size=11, color=ACCENT_TEXT, bold=True)

    title = str(lesson.get("title") or lesson.get("topic") or "课堂").strip()
    _textbox(slide, 0.9, 1.45, 11.4, 1.7, title, size=40, bold=True, line_spacing=1.15)

    summary = str(lesson.get("summary") or "").strip()
    if summary:
        _textbox(slide, 0.9, 3.25, 11.0, 0.9, summary, size=17, color=INK2, line_spacing=1.4)

    objectives = _clean_bullets(lesson.get("objectives"), 4)
    if objectives:
        _textbox(slide, 0.9, 4.55, 6.0, 0.3, "本课目标", size=11, color=INK3, bold=True)
        _paragraphs(slide, 0.9, 4.92, 11.2, 1.5, objectives, size=14, color=INK2, bullet=True)

    _textbox(slide, 0.9, 6.72, 11.5, 0.3, "Edu · 轻量级 AI 教学服务", size=9, color=INK3)
    _note(slide, str(lesson.get("summary") or ""))
    return page + 1


def _add_content(prs, lesson: dict, scene: dict, index: int, total: int, page: int) -> int:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    footer_title = str(lesson.get("title") or "课堂")
    _header(slide, f"讲解 · {index} / {total}", str(scene.get("title") or "讲解"), page, footer_title)

    bullets = _clean_bullets(scene.get("bullets"), 5)
    if bullets:
        _paragraphs(
            slide, 0.9, 2.32, 11.4, 2.9, bullets,
            size=_bullets_size(bullets), bullet=True, space_after=12,
        )

    prompt = str(scene.get("prompt") or "").strip()
    if prompt:
        _prompt_box(slide, f"课堂提问：{prompt}")

    _note(slide, str(scene.get("notes") or ""))
    return page + 1


def _add_discussion(prs, lesson: dict, scene: dict, index: int, total: int, page: int) -> int:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    footer_title = str(lesson.get("title") or "课堂")
    _header(slide, f"讨论 · {index} / {total}", str(scene.get("title") or "讨论"), page, footer_title)
    prompt = str(scene.get("prompt") or "").strip()
    _textbox(slide, 0.9, 2.45, 11.4, 1.6, prompt, size=24, color=INK, line_spacing=1.45)
    hints = _clean_bullets(scene.get("hints"), 3)
    if hints:
        _paragraphs(slide, 0.9, 4.5, 11.4, 1.3, hints, size=14, color=INK2, bullet=True)
    _textbox(slide, 0.9, 6.1, 11.4, 0.4, "想一想，再和 AI 老师继续讨论。", size=13, color=INK3)
    _note(slide, str(scene.get("notes") or prompt))
    return page + 1


def _add_quiz(prs, lesson: dict, scene: dict, index: int, total: int, page: int) -> int:
    footer_title = str(lesson.get("title") or "课堂")
    questions = [item for item in (scene.get("questions") or []) if isinstance(item, dict)]
    if not questions:
        return page
    chunks = [questions[start:start + 2] for start in range(0, len(questions), 2)]
    for chunk_index, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        suffix = "" if chunk_index == 0 else "（续）"
        _header(
            slide,
            f"检测 · {index} / {total}{suffix}",
            str(scene.get("title") or "随堂检测"),
            page,
            footer_title,
        )
        top = 2.28
        for local_index, question in enumerate(chunk, 1):
            number = chunk_index * 2 + local_index
            lines = _question_lines(question, number)
            _paragraphs(
                slide, 0.9, top, 11.4, 2.2, lines,
                size=15 if len(lines) > 1 else 16, space_after=4, line_spacing=1.25,
            )
            top += 2.15
        notes = ["答案与解析（教师用）"]
        for local_index, question in enumerate(chunk, 1):
            number = chunk_index * 2 + local_index
            notes.extend(_answer_lines(question, number))
        _note(slide, "\n".join(notes))
        page += 1
    return page


def _add_summary(prs, lesson: dict, scene: dict, index: int, total: int, page: int) -> int:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    footer_title = str(lesson.get("title") or "课堂")
    _header(slide, f"回顾 · {index} / {total}", str(scene.get("title") or "本课回顾"), page, footer_title)
    bullets = _clean_bullets(scene.get("bullets"), 6)
    if bullets:
        _paragraphs(slide, 0.9, 2.35, 11.4, 3.4, bullets, size=18, bullet=True, space_after=14)
    _prompt_box(slide, "下一步：用「测验」页做一组新题，系统会记录你的掌握度。", 5.6)
    _note(slide, str(scene.get("notes") or ""))
    return page + 1


def _add_answer_appendix(prs, lesson: dict, question_items: list[tuple[int, dict]], page: int) -> int:
    if not question_items:
        return page
    footer_title = str(lesson.get("title") or "课堂")
    chunks = [question_items[start:start + 2] for start in range(0, len(question_items), 2)]
    for chunk_index, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _header(
            slide,
            "附录 · 教师版",
            "参考答案与解析" + ("" if chunk_index == 0 else "（续）"),
            page,
            footer_title,
        )
        top = 2.3
        for number, question in chunk:
            lines = _answer_lines(question, number)
            _paragraphs(
                slide, 0.9, top, 11.4, 2.2, lines,
                size=12.5, space_after=4, line_spacing=1.3,
            )
            top += 2.2
        _note(slide, "答案附录仅供教师参考；如需隐藏，可在放映前删除本页。")
        page += 1
    return page


def build(lesson: dict) -> bytes:
    """返回 PPTX 二进制内容。"""
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    prs.core_properties.title = str(lesson.get("title") or "课堂")
    prs.core_properties.author = "Edu"
    prs.core_properties.subject = "Edu 课堂导出"

    scenes = [scene for scene in (lesson.get("scenes") or []) if isinstance(scene, dict)]
    page = 1
    answer_items: list[tuple[int, dict]] = []
    question_number = 0
    for index, scene in enumerate(scenes, 1):
        scene_type = scene.get("type") or "content"
        if scene_type == "cover":
            page = _add_cover(prs, lesson, page)
        elif scene_type == "discussion":
            page = _add_discussion(prs, lesson, scene, index, len(scenes), page)
        elif scene_type == "quiz":
            page = _add_quiz(prs, lesson, scene, index, len(scenes), page)
            for question in scene.get("questions") or []:
                if isinstance(question, dict):
                    question_number += 1
                    answer_items.append((question_number, question))
        elif scene_type == "summary":
            page = _add_summary(prs, lesson, scene, index, len(scenes), page)
        else:
            page = _add_content(prs, lesson, scene, index, len(scenes), page)

    page = _add_answer_appendix(prs, lesson, answer_items, page)

    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
