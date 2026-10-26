"""资料上传与解析。"""
from __future__ import annotations

import io
import time
from pathlib import Path

from .. import rag, store

TEXT_SUFFIXES = {".txt", ".md", ".markdown", ".csv", ".json", ".py", ".java", ".c", ".cpp", ".js", ".ts", ".html", ".css", ".xml", ".yaml", ".yml"}
PDF_SUFFIXES = {".pdf"}
DOCX_SUFFIXES = {".docx"}
PPTX_SUFFIXES = {".pptx"}
MAX_BYTES = 10 * 1024 * 1024


class MaterialError(ValueError):
    pass


def extract_text(filename: str, data: bytes, *, max_bytes: int = MAX_BYTES) -> str:
    suffix = Path(filename).suffix.lower()
    if len(data) > max_bytes:
        mb = max(1, round(max_bytes / (1024 * 1024)))
        raise MaterialError(f"文件超过 {mb} MB，请拆分后再上传。")
    if suffix in PDF_SUFFIXES:
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise MaterialError("缺少 PDF 解析依赖，请运行：pip install pypdf") from exc
        try:
            reader = PdfReader(io.BytesIO(data))
            return "\n\n".join((page.extract_text() or "") for page in reader.pages)
        except Exception as exc:  # noqa: BLE001 - 解析库异常类型不统一
            raise MaterialError(f"PDF 解析失败：{exc}") from exc
    if suffix in DOCX_SUFFIXES:
        try:
            import docx
        except ImportError as exc:
            raise MaterialError("缺少 Word 解析依赖，请运行：pip install python-docx") from exc
        try:
            document = docx.Document(io.BytesIO(data))
            parts = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except Exception as exc:  # noqa: BLE001
            raise MaterialError(f"Word 解析失败：{exc}") from exc
    if suffix in PPTX_SUFFIXES:
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise MaterialError("缺少 PPT 解析依赖，请运行：pip install python-pptx") from exc
        try:
            deck = Presentation(io.BytesIO(data))
            parts = []
            for index, slide in enumerate(deck.slides, 1):
                texts = [
                    shape.text.strip()
                    for shape in slide.shapes
                    if getattr(shape, "has_text_frame", False) and shape.text.strip()
                ]
                if texts:
                    parts.append(f"第 {index} 页：\n" + "\n".join(texts))
            return "\n\n".join(parts)
        except Exception as exc:  # noqa: BLE001
            raise MaterialError(f"PPT 解析失败：{exc}") from exc
    if suffix in TEXT_SUFFIXES or suffix == "":
        for encoding in ("utf-8-sig", "utf-8", "gb18030", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
        raise MaterialError("无法识别文件编码，请另存为 UTF-8 后重试。")
    raise MaterialError(f"暂不支持 {suffix} 格式，请上传 txt / md / pdf / docx / pptx。")


def create(filename: str, data: bytes) -> dict:
    text = extract_text(filename, data).strip()
    if len(text) < 20:
        raise MaterialError("没有提取到足够的文字内容（可能扫描件没有文字层）。")
    text = text[:400_000]
    chunks = rag.chunk_text(text)
    material = {
        "id": store.new_id("mat_"),
        "title": Path(filename).stem[:80] or "未命名资料",
        "filename": filename[:160],
        "chars": len(text),
        "chunks": [{"id": f"c{index + 1}", "text": chunk} for index, chunk in enumerate(chunks)],
        "preview": rag.preview(text, 220),
        "created_at": time.time(),
    }
    store.save_item("materials", material)
    return material


def list_materials() -> list[dict]:
    items = store.list_items("materials")
    return [
        {
            "id": item.get("id"),
            "title": item.get("title"),
            "filename": item.get("filename"),
            "chars": item.get("chars", 0),
            "chunks": len(item.get("chunks") or []),
            "preview": item.get("preview", ""),
            "created_at": item.get("created_at"),
        }
        for item in items
    ]
